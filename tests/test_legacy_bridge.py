"""Synthetic-only validation of lossless V2/V3 archive and conflict-safe Drax bridge."""
import os
import sqlite3
from pathlib import Path
import pytest
from openpyxl import Workbook, load_workbook

import v31_legacy_history as legacy
import v31_drax_bridge as bridge


def make_old_db(path):
    with sqlite3.connect(path) as c:
        c.executescript("""
        CREATE TABLE shifts (id INTEGER PRIMARY KEY,date TEXT NOT NULL,
            shift_type TEXT NOT NULL,lead_name TEXT NOT NULL,created_at TEXT NOT NULL,
            is_active INTEGER NOT NULL,opened_by_assoc TEXT NOT NULL DEFAULT '');
        CREATE TABLE entries (id INTEGER PRIMARY KEY,shift_id INTEGER,tech_name TEXT NOT NULL);
        CREATE TABLE reports (id INTEGER PRIMARY KEY,shift_id INTEGER,content TEXT NOT NULL);
        CREATE TABLE associates (assoc_num TEXT PRIMARY KEY,full_name TEXT,display_name TEXT,
            default_id TEXT,shift TEXT,status TEXT,is_manager INTEGER,is_mst INTEGER);
        CREATE TABLE auxiliary (id INTEGER PRIMARY KEY,value TEXT NOT NULL);
        INSERT INTO shifts VALUES (15,'2026-10-24','S4','Synthetic A','2026-10-24 05:28',0,'111111');
        INSERT INTO shifts VALUES (18,'2026-10-25','S5','Synthetic B','2026-10-25 17:30',0,'222222');
        INSERT INTO entries VALUES (501,15,'Synthetic A');
        INSERT INTO reports VALUES (300,18,'IND3 S5 historical passdown');
        INSERT INTO auxiliary VALUES(4,'must preserve all ancillary data');
        INSERT INTO associates VALUES('111111','SYNTHETIC A','Synthetic A','a1','S1','MST',0,1);
        INSERT INTO associates VALUES('222222','SYNTHETIC B','Synthetic B','b2','S2','MST',0,1);
        """)


def make_drax(path):
    book=Workbook();sheet=book.active
    sheet.append(["Name","Assoc. #","Default ID","Shift","Job Title","Other Formula"])
    sheet.append(["SYNTHETIC A",111111,"a1","S1","MST","=2+2"])
    sheet.append(["SYNTHETIC B",222222,"b2","S2","MST","keep"])
    book.save(path)


def test_original_v2_v3_sqlite_exact_history_and_source_read_only(tmp_path):
    src=tmp_path/"legacy.db";make_old_db(src)
    before=src.read_bytes()
    result=legacy.import_snapshot(src,tmp_path/"private")
    assert result["shifts"]==2 and result["reports"]==1 and result["entries"]==1
    archive=tmp_path/"private"/"legacy-v2-v3-archive.sqlite3"
    with sqlite3.connect(src) as o, sqlite3.connect(archive) as a:
        original_schema={r[0]:r[1] for r in o.execute(
            "SELECT name,sql FROM sqlite_master WHERE type='table'")}
        new_schema={r[0]:r[1] for r in a.execute(
            "SELECT name,sql FROM sqlite_master WHERE type='table'")}
        assert original_schema==new_schema
        assert a.execute("SELECT value FROM auxiliary WHERE id=4").fetchone()[0]=="must preserve all ancillary data"
        assert a.execute("SELECT id FROM shifts ORDER BY id").fetchall()==[(15,),(18,)]
    assert before==src.read_bytes()
    manager=legacy.history(tmp_path/"private",None)
    assert manager["total"]==2 and manager["rows"][0]["id"]==18
    assert manager["rows"][0]["report_count"]==1
    mst=legacy.history(tmp_path/"private","111111")
    assert mst["total"]==1 and mst["rows"][0]["id"]==15


def test_legacy_import_rejects_unrecognized_database(tmp_path):
    source=tmp_path/"bad.db"
    with sqlite3.connect(source) as c:
        c.execute("CREATE TABLE unrelated(x)")
    with pytest.raises(legacy.LegacyHistoryError):
        legacy.inspect(source)
    assert not (tmp_path/"legacy-v2-v3-archive.sqlite3").exists()


def test_drax_bridge_baseline_does_not_change_original_files(tmp_path,monkeypatch):
    old=tmp_path/"old.db";make_old_db(old)
    xlsx=tmp_path/"Drax  Associates.xlsx";make_drax(xlsx)
    old_bytes=old.read_bytes();drax_bytes=xlsx.read_bytes()
    monkeypatch.setenv("IND3_V31_DRAX_WRITE_ENABLED","1")
    r=bridge.reconcile(old,xlsx,tmp_path/"private",write=True)
    assert r["first_baseline"] and not r["workbook_updated"]
    assert old.read_bytes()==old_bytes
    assert xlsx.read_bytes()==drax_bytes


def test_old_v3_edit_reaches_shared_drax_if_explicitly_enabled(tmp_path,monkeypatch):
    old=tmp_path/"old.db";make_old_db(old)
    xlsx=tmp_path/"Drax  Associates.xlsx";make_drax(xlsx)
    private=tmp_path/"private"
    bridge.reconcile(old,xlsx,private)
    with sqlite3.connect(old) as c:
        c.execute("UPDATE associates SET shift='S3', full_name='UPDATED NAME' WHERE assoc_num='111111'")
    source_before=old.read_bytes()
    monkeypatch.delenv("IND3_V31_DRAX_WRITE_ENABLED",raising=False)
    preview=bridge.reconcile(old,xlsx,private,write=True)
    assert not preview["write_enabled"] and preview["candidate_cell_changes"]==2
    assert load_workbook(xlsx).active["D2"].value=="S1"
    monkeypatch.setenv("IND3_V31_DRAX_WRITE_ENABLED","1")
    done=bridge.reconcile(old,xlsx,private,write=True)
    assert done["workbook_updated"] and done["candidate_cell_changes"]==2
    book=load_workbook(xlsx)
    assert book.active["D2"].value=="S3"
    assert book.active["A2"].value=="UPDATED NAME"
    assert book.active["F2"].value=="=2+2"
    assert Path(done["backup_created"]).is_file()
    assert old.read_bytes()==source_before
    again=bridge.reconcile(old,xlsx,private,write=True)
    assert not again["workbook_updated"]


def test_independent_drax_edits_are_not_overwritten(tmp_path,monkeypatch):
    old=tmp_path/"old.db";make_old_db(old)
    xlsx=tmp_path/"Drax  Associates.xlsx";make_drax(xlsx)
    private=tmp_path/"private"
    bridge.reconcile(old,xlsx,private)
    book=load_workbook(xlsx)
    book.active["D2"]="S5";book.save(xlsx)
    monkeypatch.setenv("IND3_V31_DRAX_WRITE_ENABLED","1")
    result=bridge.reconcile(old,xlsx,private,write=True)
    assert not result["workbook_updated"]
    assert load_workbook(xlsx).active["D2"].value=="S5"


def test_simultaneous_conflicting_edits_are_reported_not_overwritten(tmp_path,monkeypatch):
    old=tmp_path/"old.db";make_old_db(old)
    xlsx=tmp_path/"Drax  Associates.xlsx";make_drax(xlsx)
    private=tmp_path/"private"
    bridge.reconcile(old,xlsx,private)
    with sqlite3.connect(old) as c:
        c.execute("UPDATE associates SET shift='S3' WHERE assoc_num='111111'")
    book=load_workbook(xlsx);book.active["D2"]="S4";book.save(xlsx)
    monkeypatch.setenv("IND3_V31_DRAX_WRITE_ENABLED","1")
    result=bridge.reconcile(old,xlsx,private,write=True)
    assert result["conflicts"] and not result["workbook_updated"]
    assert load_workbook(xlsx).active["D2"].value=="S4"


def test_v3_deletion_is_reported_for_manual_review(tmp_path):
    old=tmp_path/"old.db";make_old_db(old)
    xlsx=tmp_path/"Drax  Associates.xlsx";make_drax(xlsx)
    private=tmp_path/"private"
    bridge.reconcile(old,xlsx,private)
    with sqlite3.connect(old) as c:
        c.execute("DELETE FROM associates WHERE assoc_num='222222'")
    outcome=bridge.reconcile(old,xlsx,private)
    assert any(x["assoc"]=="222222" and x["field"]=="row" for x in outcome["conflicts"])
    assert load_workbook(xlsx).active["B3"].value==222222
